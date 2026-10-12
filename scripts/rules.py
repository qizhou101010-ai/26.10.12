"""审核规则：读取 审核规则.xlsx，并执行不需要大模型的硬性检查。

审核规则.xlsx 有三张表，用户可以直接在 Excel 里修改：
  列设置   —— 填报表里每种信息对应哪一列，可写多个别名（列名不同时改这里）
  项目规则 —— 每个填报项目的必填、类型、单位、范围、小数位、是否核查出处、补充要求
  通用规则 —— 适用于所有行的文字规则，交给大模型一起判断

填报表里每行自带的“字符要求”“数据要求”（如“纯数字，6位”“保留1位小数”）会被自动读成规则；
项目规则 里明确填写的项优先。
"""
import re
from dataclasses import dataclass, field
from datetime import date, datetime

from openpyxl import load_workbook

import config

# 角色 -> 可能的列名（按顺序匹配表头，命中第一个）。“字符”列可以没有。
DEFAULT_COLUMNS = {
    "项目": ["数据类型", "填报项目", "数据项", "指标名称", "指标", "项目名称", "字段名称", "字段"],
    "要求": ["数据要求", "填写要求", "填写说明", "填报说明", "说明"],
    "字符": ["字符要求", "格式要求", "格式"],
    "内容": ["数据内容", "填写内容", "填报值", "填报内容", "填报数据", "数值", "值"],
    "来源": ["数据来源", "资料来源", "出处", "资料出处", "来源"],
    "日期": ["来源日期", "来源发布日期", "出处日期", "资料日期", "发布日期"],
}
RULE_FIELDS = ["填报项目", "必填", "数据类型", "字符类型", "字符数", "单位", "取值范围", "小数位",
               "需核查出处", "补充要求", "正则表达式"]
DATE_RE = re.compile(r"^\d{4}-\d{2}(-\d{2})?$")
LOOSE_DATE_RE = re.compile(r"^\s*(\d{4})\s*[-./年]\s*(\d{1,2})\s*(?:月\s*)?(?:[-./]?\s*(\d{1,2})\s*日?)?\s*$")


def normalize_date(s: str) -> str | None:
    """把“2018-6”“2018/6/5”“2018.06”“2018年6月”等写法统一成 YYYY-MM 或 YYYY-MM-DD；认不出返回 None。"""
    m = LOOSE_DATE_RE.match(s or "")
    if not m or not 1 <= int(m[2]) <= 12 or (m[3] and not 1 <= int(m[3]) <= 31):
        return None
    return f"{m[1]}-{int(m[2]):02d}" + (f"-{int(m[3]):02d}" if m[3] else "")

# 数据类型列为这些值的行，其内容是整张表（或下一段）描述的对象名称，后续各行都针对该对象核查
SUBJECT_ITEMS = ("名称", "对象名称", "名字")


_DEG = r"(\d+(?:\.\d+)?)\s*°?\s*(?:(\d+(?:\.\d+)?)\s*[′']\s*)?(?:(\d+(?:\.\d+)?)\s*[″\"]\s*)?"


def _to_degrees(m) -> float:
    d, mi, se = (float(x) if x else 0.0 for x in m.groups())
    return d + mi / 60 + se / 3600


def parse_coordinate(s: str) -> tuple[float, float] | None:
    """把各种写法的经纬度换算成 (经度, 纬度) 十进制度；西经、南纬为负。认不出返回 None。"""
    s = s.replace("：", "").replace(":", "")
    lon = re.search(r"(东经|西经|[EW])\s*" + _DEG, s)
    lat = re.search(r"(北纬|南纬|[NS])\s*" + _DEG, s)
    if lon and lat:
        sign = lambda m: -1 if m.group(1) in ("西经", "W", "南纬", "S") else 1
        return (sign(lon) * _to_degrees(re.match(_DEG, s[lon.start(2):])),
                sign(lat) * _to_degrees(re.match(_DEG, s[lat.start(2):])))
    m = re.fullmatch(r"\s*(-?\d{1,3}(?:\.\d+)?)°?\s*[,，、]\s*(-?\d{1,3}(?:\.\d+)?)°?\s*", s)
    return (float(m[1]), float(m[2])) if m else None


def compare_coordinates(filled: str, source_text: str, tolerance: float = 1 / 60) -> str | None:
    """返回 "一致" / "经纬度写反" / "不一致"；任一方解析不出返回 None（交给大模型判断）。"""
    a, b = parse_coordinate(filled), parse_coordinate(source_text)
    if not a or not b:
        return None
    close = lambda p, q: abs(p[0] - q[0]) <= tolerance and abs(p[1] - q[1]) <= tolerance
    if close(a, b):
        return "一致"
    if close((a[1], a[0]), b):
        return "经纬度写反"
    return "不一致"


def is_coordinate(s: str) -> bool:
    """经纬度：“东经102°30′15″，北纬25°10′05″”“E102.5 N25.2”或“102.5042, 25.1681”。"""
    if re.fullmatch(r"-?\d{1,3}(\.\d+)?°?\s*[,，、]\s*-?\d{1,3}(\.\d+)?°?", s):
        return True
    return (bool(re.search(r"东经|西经|[EW]", s)) and bool(re.search(r"北纬|南纬|[NS]", s))
            and len(re.findall(r"\d+(?:\.\d+)?", s)) >= 2)


# 字符类型 -> 必须整体匹配的正则（或判断函数）
CHAR_TYPES = {
    "纯汉字": r"[一-鿿·]+",
    "纯数字": r"\d+",
    "纯字母": r"[A-Za-z]+",
    "字母和数字": r"[A-Za-z0-9]+",
    "手机号": r"1[3-9]\d{9}",
    "邮箱": r"[\w.+-]+@[\w-]+(\.[\w-]+)+",
    "身份证号": r"\d{17}[\dXx]",
    "统一社会信用代码": r"[0-9A-HJ-NPQRTUWXY]{2}\d{6}[0-9A-HJ-NPQRTUWXY]{10}",
    "邮政编码": r"\d{6}",
    "经纬度": is_coordinate,
}


def parse_length(spec: str) -> tuple[int, float] | None:
    """字符数写法：6 / 2~10 / ≤50 / ≥2（也认 <= >= 最多 至少）。"""
    spec = spec.replace(" ", "")
    if m := re.fullmatch(r"(\d+)", spec):
        return int(m[1]), int(m[1])
    if m := re.fullmatch(r"(\d+)[~～\-至到](\d+)", spec):
        return int(m[1]), int(m[2])
    if m := re.fullmatch(r"(?:≤|<=|最多|不超过)(\d+)", spec):
        return 0, int(m[1])
    if m := re.fullmatch(r"(?:≥|>=|至少|不少于)(\d+)", spec):
        return int(m[1]), float("inf")
    return None


def describe_length(lo: int, hi: float) -> str:
    if lo == hi:
        return f"正好 {lo} 个字符"
    if hi == float("inf"):
        return f"至少 {lo} 个字符"
    return f"{lo}~{hi} 个字符" if lo else f"最多 {hi} 个字符"


@dataclass
class Issue:
    字段: str              # 出问题的列：内容 / 来源 / 日期（对应 列设置 里的角色）
    问题类型: str
    说明: str
    建议值: str | None = None
    证据: str = ""
    状态: str = "错误"      # 错误：有把握的问题；需复核：需要人工确认


@dataclass
class Rules:
    columns: dict = field(default_factory=lambda: {k: list(v) for k, v in DEFAULT_COLUMNS.items()})
    items: dict = field(default_factory=dict)      # 填报项目 -> 规则字典
    general: list = field(default_factory=list)


def cell_str(v) -> str:
    """把单元格的值转成文本；日期统一成 YYYY-MM-DD，整数不带 .0。"""
    if v is None:
        return ""
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def load_rules(path=config.RULES_FILE) -> Rules:
    rules = Rules()
    if not path.exists():
        return rules
    wb = load_workbook(path, read_only=True, data_only=True)
    if "列设置" in wb.sheetnames:
        for row in list(wb["列设置"].iter_rows(values_only=True))[1:]:
            if row and row[0] and row[1]:
                rules.columns[cell_str(row[0])] = [n for n in re.split(r"[,，、\s]+", cell_str(row[1])) if n]
    if "项目规则" in wb.sheetnames:
        rows = list(wb["项目规则"].iter_rows(values_only=True))
        header = [cell_str(h) for h in rows[0]] if rows else []
        for row in rows[1:]:
            item = {h: cell_str(v) for h, v in zip(header, row)}
            if item.get("填报项目"):
                rules.items[item["填报项目"]] = item
    if "通用规则" in wb.sheetnames:
        rules.general = [cell_str(r[0]) for r in list(wb["通用规则"].iter_rows(values_only=True))[1:]
                         if r and r[0]]
    return rules


# 字符要求里的说法 -> 字符类型（长的、组合的说法放前面）
CHAR_TYPE_WORDS = [
    ("经纬度", "经纬度"), ("坐标", "经纬度"),
    ("统一社会信用代码", "统一社会信用代码"), ("身份证", "身份证号"), ("手机", "手机号"),
    ("邮箱", "邮箱"), ("邮编", "邮政编码"), ("邮政编码", "邮政编码"),
    ("字母和数字", "字母和数字"), ("字母数字", "字母和数字"), ("数字和字母", "字母和数字"),
    ("数字或字母", "字母和数字"), ("字母或数字", "字母和数字"),
    ("纯汉字", "纯汉字"), ("汉字", "纯汉字"), ("中文", "纯汉字"),
    ("纯数字", "纯数字"), ("数字", "纯数字"), ("纯字母", "纯字母"), ("英文字母", "纯字母"), ("字母", "纯字母"),
    ("不限", "不限"),
]
TYPE_WORDS = {"数值": "数值", "数字": "数值", "文本": "文本", "字符": "文本", "日期": "日期"}


def parse_row_requirements(row: dict) -> dict:
    """把填报表里这一行自带的“字符要求 / 数据要求”读成结构化规则。读不出的部分交给大模型参考。"""
    rule = {}
    char_req = row.get("字符", "")
    data_req = row.get("要求", "")
    if row.get("项目") in TYPE_WORDS:          # “数据类型”列写的是 数值/文本/日期
        rule["数据类型"] = TYPE_WORDS[row["项目"]]
    if char_req:
        for word, char_type in CHAR_TYPE_WORDS:
            if word in char_req:
                rule["字符类型"] = char_type
                break
        n = r"(\d+)\s*(?:位|个|字)"
        if m := re.search(r"(\d+)\s*[~～\-至到]\s*" + n, char_req):
            rule["字符数"] = f"{m[1]}~{m[2]}"
        elif m := re.search(r"(?:不超过|最多|不多于|≤|<=)\s*" + n, char_req):
            rule["字符数"] = f"≤{m[1]}"
        elif m := re.search(r"(?:不少于|至少|≥|>=)\s*" + n, char_req):
            rule["字符数"] = f"≥{m[1]}"
        elif m := re.search(n, char_req):
            rule["字符数"] = m[1]
    if m := re.search(r"保留\s*(\d+)\s*位小数", data_req):
        rule["小数位"] = m[1]
    elif "整数" in data_req:
        rule["小数位"] = "0"
    if rule.get("字符类型") == "纯数字" and "数据类型" not in rule:
        # 有单位或小数要求的是数值（允许小数点）；否则当编号类文本（如代码），按字符逐个检查
        is_measure = "小数位" in rule or "单位" in data_req
        rule["数据类型"] = "数值" if is_measure else "文本"
    if re.search(r"不需要?(注明)?(资料|数据)?来源|无需(注明)?来源", data_req + char_req):
        rule["需核查出处"] = "否"
    return rule


def effective_rule(row: dict, rules: Rules) -> dict:
    """本行最终规则 = 表格里每行自带的要求 + 审核规则.xlsx 中该项目明确填写的项（后者优先）。"""
    configured = {k: v for k, v in rules.items.get(row["项目"], {}).items() if v}
    return {**parse_row_requirements(row), **configured}


def to_number(s: str) -> float | None:
    try:
        return float(s.replace(",", ""))
    except ValueError:
        return None


def check_text_format(content: str, rule: dict) -> list[Issue]:
    """字符类型、字符数、正则表达式。规则表里写错的规则只打印提示，不算填报人的错。"""
    issues = []
    char_type = rule.get("字符类型", "")
    if char_type and char_type != "不限":
        pattern = CHAR_TYPES.get(char_type)
        if char_type == "纯数字" and rule.get("数据类型") == "数值":
            pass                                   # 数值项允许小数点，由数值检查负责
        elif pattern is None:
            print(f"  提示：项目“{rule.get('填报项目')}”的字符类型“{char_type}”无法识别，已跳过该检查")
        elif not (pattern(content) if callable(pattern) else re.fullmatch(pattern, content)):
            issues.append(Issue("内容", "字符类型错误", f"应为{char_type}，实际填写“{content}”"))

    length_spec = rule.get("字符数", "")
    if length_spec:
        bounds = parse_length(length_spec)
        if bounds is None:
            print(f"  提示：项目“{rule.get('填报项目')}”的字符数“{length_spec}”无法识别，已跳过该检查")
        elif not (bounds[0] <= len(content) <= bounds[1]):
            issues.append(Issue("内容", "字符数不符",
                                f"应为{describe_length(*bounds)}，实际 {len(content)} 个"))

    regex = rule.get("正则表达式", "")
    if regex:
        try:
            if not re.fullmatch(regex, content):
                issues.append(Issue("内容", "格式错误", f"不符合规定格式（{regex}）"))
        except re.error:
            print(f"  提示：项目“{rule.get('填报项目')}”的正则表达式写法有误，已跳过该检查")
    return issues


def hard_check(row: dict, rule: dict) -> list[Issue]:
    """row 的键是角色名：项目/要求/内容/来源/日期。"""
    issues = []
    content, source, pub_date = row["内容"], row["来源"], row["日期"]
    need_source = rule.get("需核查出处", "是") == "是"

    if rule.get("必填", "是") == "是" and not content:
        issues.append(Issue("内容", "漏填", "该项为必填，但没有填写内容"))
    if content and rule.get("数据类型") == "数值":
        num = to_number(content)
        if num is None:
            issues.append(Issue("内容", "格式错误",
                                f"应填写纯数字（单位：{rule.get('单位') or '见要求'}），不能带文字或单位"))
        else:
            rng = rule.get("取值范围", "")
            m = re.fullmatch(r"\s*([-\d.]+)\s*[~～-]\s*([-\d.]+)\s*", rng)
            if m and not (float(m.group(1)) <= num <= float(m.group(2))):
                issues.append(Issue("内容", "超出合理范围",
                                    f"数值 {content} 不在合理范围 {rng}（{rule.get('单位', '')}）内，"
                                    "可能是单位或数量级错误", 状态="需复核"))
            digits = rule.get("小数位", "")
            if digits.isdigit():
                actual = len(content.split(".")[1]) if "." in content else 0
                if actual != int(digits):
                    issues.append(Issue("内容", "格式错误", f"应保留 {digits} 位小数"))
    if content:
        issues += check_text_format(content, rule)
    if need_source:
        if not source:
            issues.append(Issue("来源", "漏填", "该项需要注明来源"))
        elif not re.fullmatch(r"《[^《》]+》", source):
            issues.append(Issue("来源", "格式错误", "来源须用书名号《》填写资料全称",
                                建议值=f"《{source.strip('《》<> ')}》"))
        if not pub_date:
            issues.append(Issue("日期", "漏填", "该项需要填写来源日期"))
        elif not DATE_RE.match(pub_date):
            fixed = normalize_date(pub_date)
            issues.append(Issue("日期", "格式错误", "日期格式应为 YYYY-MM-DD 或 YYYY-MM" +
                                (f"，“{pub_date}”应写成 {fixed}" if fixed else ""), 建议值=fixed))
    return issues
