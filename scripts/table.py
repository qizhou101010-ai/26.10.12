"""生成标准格式的填报表（表头：数据类型 / 数据要求 / 字符要求 / 数据内容 / 数据来源 / 来源日期）。"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

HEADER = ["数据类型", "数据要求", "字符要求", "数据内容", "数据来源", "来源日期"]
DEFAULT_ITEMS = ["名称", "高度", "地理坐标", "地形特点", "覆盖植被类型"]

# 不用大模型时，常见数据类型的默认 数据要求 / 字符要求（写法能被程序自动读懂）
DEFAULT_REQUIREMENTS = {
    "名称": ("对象的规范名称，与资料中的写法一致", "纯汉字，2~20个字"),
    "高度": ("最高峰（主峰）海拔，单位：米，填整数，以资料中的数值为准", "纯数字"),
    "海拔": ("最高峰（主峰）海拔，单位：米，填整数，以资料中的数值为准", "纯数字"),
    "地理坐标": ("最高峰（主峰）位置，经度在前、纬度在后，度分秒或十进制度均可", "经纬度"),
    "地形特点": ("简要描述山体走向、坡度和地貌类型，使用资料中的说法，不添加资料没有的内容", "不超过100字"),
    "覆盖植被类型": ("列出资料中提到的主要植被类型，用顿号分隔", "不超过50字"),
    "面积": ("单位：平方公里，保留1位小数", "数字"),
    "长度": ("单位：公里，保留1位小数", "数字"),
    "全长": ("单位：公里，保留1位小数", "数字"),
    "人口": ("单位：万人，保留1位小数，注明统计年份", "数字"),
}


def write_table(path, title: str, rows: list[list]):
    """rows 每行按 HEADER 的顺序给值，缺的列留空。第一行是合并的标题。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "填报表"
    ws.append([title])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEADER))
    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.append(HEADER)
    for c in ws[2]:
        c.font = Font(bold=True)
    for row in rows:
        ws.append(list(row) + [None] * (len(HEADER) - len(row)))
    for col, w in zip("ABCDEF", [14, 40, 22, 40, 36, 12]):
        ws.column_dimensions[col].width = w
    for cells in ws.iter_rows(min_row=3):
        for c in cells:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
