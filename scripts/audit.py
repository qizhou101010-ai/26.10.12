"""批量审核：逐行检查 待审核/ 里的所有 Excel，输出错误报告和修正版。

用法：
    python audit.py                  # 审核 待审核/ 下所有 .xlsx
    python audit.py 某个文件.xlsx     # 只审核指定文件
"""
import sys
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill

import config
import program_check
from kb import KnowledgeBase, evidence_in_doc
from llm import ask_json
from rules import (SUBJECT_ITEMS, Issue, Rules, cell_str, compare_coordinates, effective_rule, hard_check,
                   load_rules, to_number)

FIXED_FILL = PatternFill("solid", fgColor="FFF2CC")    # 黄色：已自动修正
REVIEW_FILL = PatternFill("solid", fgColor="F8CBAD")   # 橙色：需人工复核
LONG_TEXT = 30   # 超过这个字数的内容不自动改写，只给建议


def is_descriptive(rule: dict, content: str) -> bool:
    """描述性内容（地形特点、植被类型等）：不是数值、也没规定字符类型，或者很长。这类只给建议不自动改。"""
    plain = rule.get("数据类型") != "数值" and rule.get("字符类型", "不限") in ("", "不限")
    return plain or len(content) > LONG_TEXT

SYSTEM_PROMPT = """你是严谨的数据审核员，负责核查填报表中的数据是否有资料依据。
规则：
1. 只能依据用户提供的【资料原文片段】判断，不得使用你自己的知识补充或推测。
2. 注意单位换算、数量级、统计年份/口径是否一致。
3. 引用原文时必须逐字摘录，不能改写。
4. 资料说的必须是【核查对象】本身，不能拿其他对象的数据来证明。
5. 描述性文字（如地形特点、植被类型）：意思与原文一致即为支持，不要求逐字相同；
   含有原文没有的说法为“部分支持”，与原文矛盾为“不支持”，并在内容说明中逐条列出这些说法。
6. 地理坐标：度分秒与十进制度可互相换算（1°=60′，1′=60″），换算后相差不超过1′视为一致；注意经度纬度是否写反。
7. 原文中被引用、转述后又被否定、辟谣或质疑的说法，不能算作支持；要看资料作者自己的结论。
   如果资料明确驳斥了该说法，应判为“不支持”。
8. 以 JSON 格式输出。"""

USER_PROMPT = """【核查对象】{主体}
【填报项目】{项目}（是核查对象的一个属性）
【数据要求】{要求}
【字符要求】{字符}
【该项规则】{rule_text}
【通用规则】{general}

【填写内容】{内容}
【资料来源】{来源}
【来源发布日期（填报人填写）】{日期}

【资料文件】{doc_name}
【资料原文片段】
{passages}

请完成：
1. 从原文片段中找出该资料的发布/成文日期。判断填报人填写的来源发布日期是否一致：只填到月份时年月一致即算一致；找不到日期时 日期一致 填 null。
2. 判断资料原文能否证明【填写内容】：支持 / 不支持（原文有该数据但与填写不同）/ 部分支持 / 未提及（原文没有该数据）。
3. 若为“不支持”且原文中有正确数据，按填写要求的单位和格式给出建议填写内容（只填值本身）。
4. 只看【填写内容】这个值本身是否违反填写要求或规则（如单位、小数位、是否只填数字），没有则填 null。不要重复第1、2步已发现的日期或内容问题。

输出 JSON，字段如下：
{{"资料发布日期": "YYYY-MM-DD 或 YYYY-MM 或 null",
  "日期一致": true/false/null,
  "日期说明": "一句话",
  "日期原文": "逐字摘录原文中标注日期的那句话，没有则为空字符串",
  "内容结论": "支持/不支持/部分支持/未提及",
  "原文证据": "逐字摘录能说明问题的原文，不超过100字，没有则为空字符串",
  "内容说明": "一句话说明判断理由",
  "建议填写内容": "字符串或 null",
  "规则问题": "字符串或 null",
  "置信度": "高/中/低"}}"""


def read_rows(path: Path, rules: Rules):
    """找到表头行，返回 (列号映射, 数据行列表, 表头行号)。每个数据行是 {角色: 值, '_row': 行号}。"""
    ws = load_workbook(path, data_only=True).active
    for header_row in range(1, 11):
        names = {cell_str(c.value): c.column for c in ws[header_row]}
        col_of = {role: next((names[a] for a in aliases if a in names), None)
                  for role, aliases in rules.columns.items()}
        if col_of.get("项目") and col_of.get("内容"):
            break
    else:
        raise ValueError(f"{path.name}：前 10 行里找不到表头。需要有这两类列："
                         f"{'/'.join(rules.columns['项目'])}，以及 {'/'.join(rules.columns['内容'])}。"
                         "列名不同时，请在 审核规则.xlsx 的“列设置”里添加。")
    rows, subject = [], ""
    for r in range(header_row + 1, ws.max_row + 1):
        row = {role: cell_str(ws.cell(r, c).value) if c else "" for role, c in col_of.items()}
        if row["项目"]:
            if row["项目"] in SUBJECT_ITEMS:      # 遇到“名称”行，开始一个新对象
                subject = row["内容"]
            row["主体"] = subject
            row["_row"] = r
            rows.append(row)
    return col_of, rows, header_row


def rule_text(rule: dict) -> str:
    parts = [f"{k}：{v}" for k, v in rule.items() if v and k != "填报项目"]
    return "；".join(parts) or "（无，按填写要求判断）"


def llm_available() -> bool:
    return config.USE_LLM != "off" and bool(config.LLM_API_KEY)


def program_issues(row: dict, rule: dict, content: program_check.Finding,
                   date: program_check.Finding | None) -> list[Issue]:
    """把程序核查中有把握的结论转成问题清单。"""
    issues = []
    if content.结论 == "不支持":
        auto = bool(content.建议值) and not is_descriptive(rule, row["内容"])
        issues.append(Issue("内容", "内容与资料不符", "【程序核查】" + content.说明, 建议值=content.建议值,
                            证据=content.证据, 状态="错误" if auto else "需复核"))
    elif content.结论 == "年份不符":
        issues.append(Issue("内容", "统计年份不符", "【程序核查】" + content.说明, 证据=content.证据, 状态="需复核"))
    if date and date.结论 == "不支持":
        issues.append(Issue("日期", "来源日期错误", "【程序核查】" + date.说明, 建议值=date.建议值,
                            证据=date.证据, 状态="错误"))
    return issues


def find_in_other_docs(row: dict, rule: dict, kb: KnowledgeBase, cited) -> Issue | None:
    """引用的资料里找不到时，到资料库其他资料里找：别的资料能证明，多半是出处引错了。"""
    for other in kb.docs:
        if other is cited:
            continue
        found = program_check.check_content(row, rule, other.text)
        if found.结论 == "支持":
            date, _ = program_check.find_publish_date(other.text)
            date_text = f"，该资料发布日期为 {program_check.format_date(date)}" if date else ""
            return Issue("来源", "出处可能引错",
                         f"【程序核查】《{cited.path.stem}》中找不到该内容，但《{other.path.stem}》中有{date_text}。"
                         "请确认出处并相应修改来源日期", 建议值=f"《{other.path.stem}》", 证据=found.证据, 状态="需复核")
    return None


def verify_with_source(row: dict, rule: dict, rules: Rules, kb: KnowledgeBase,
                       known: list[Issue]) -> list[Issue]:
    """程序优先：程序能确定的直接出结论；拿不准的交给模型；没有模型时交给人工。"""
    doc = kb.find(row["来源"])
    if doc is None:
        return [Issue("来源", "资料库中找不到该资料",
                      f"资料库里没有名为 {row['来源']} 的资料，无法核查内容和日期。"
                      "请确认资料名称是否写对，或把该资料导入资料库。", 状态="需复核")]

    content = program_check.check_content(row, rule, doc.text)
    if content.结论 == "不确定" and (other := find_in_other_docs(row, rule, kb, doc)):
        return [other]
    date = program_check.check_date(row["日期"], doc.text) if row["日期"] else None
    issues = program_issues(row, rule, content, date)
    content_unsure = content.结论 == "不确定"
    date_unsure = date is not None and date.结论 == "不确定"
    if not (content_unsure or date_unsure):
        return issues

    note = ""
    if llm_available():
        try:
            for issue in llm_verify(row, rule, rules, kb, doc, known + issues):
                is_date = issue.问题类型 == "来源日期错误"
                if (date_unsure and is_date) or (content_unsure and not is_date):
                    issue.说明 = "【模型判断】" + issue.说明
                    issues.append(issue)
            return issues
        except Exception as e:   # 模型调用失败（如余额不足、断网）时退回人工复核
            note = f"（大模型调用失败，已转人工：{str(e)[:60]}）"

    # 不用模型：程序拿不准的交给人工，附上相关原文方便对照
    related = content.证据 or program_check.related_text(row, doc.text)
    if content_unsure:
        issues.append(Issue("内容", "需人工核对", f"【程序核查】{content.说明}，请对照原文人工判断{note}",
                            证据=related, 状态="需复核"))
    if date_unsure:
        issues.append(Issue("日期", "需人工核对", f"【程序核查】{date.说明}，请人工确认来源日期{note}",
                            状态="需复核"))
    return issues


def llm_verify(row: dict, rule: dict, rules: Rules, kb: KnowledgeBase, doc, known: list[Issue]) -> list[Issue]:
    """交给大模型判断（只在程序拿不准时调用）。"""
    query = f"{row['主体']} {row['项目']} {row['要求']} {row['内容']}"
    passages = "\n\n".join(f"[片段{i + 1}]\n{p}" for i, p in enumerate(kb.search(doc, query)))
    result = ask_json(SYSTEM_PROMPT, USER_PROMPT.format(
        **row, rule_text=rule_text(rule), general="；".join(rules.general) or "无",
        doc_name=doc.path.name, passages=passages))

    # 模型可能把数值返回成数字类型，或返回字符串 "null"，统一成文本
    for key in ("建议填写内容", "资料发布日期"):
        value = cell_str(result.get(key))
        result[key] = None if value in ("", "null", "None") else value

    issues = []
    evidence = (result.get("原文证据") or "").strip()
    confidence = result.get("置信度", "中")
    evidence_ok = not evidence or evidence_in_doc(evidence, doc)
    sure = evidence_ok and confidence != "低"
    note = "" if evidence_ok else "（模型给出的证据在原文中找不到，请人工复核）"

    if result.get("日期一致") is False:
        date_evidence = (result.get("日期原文") or "").strip()
        date_ok = bool(date_evidence) and evidence_in_doc(date_evidence, doc) and confidence != "低"
        issues.append(Issue("日期", "来源日期错误",
                            (result.get("日期说明") or "") +
                            ("" if date_ok else "（日期原文未能在资料中核实，请人工复核）"),
                            建议值=result.get("资料发布日期"), 证据=date_evidence,
                            状态="错误" if date_ok else "需复核"))

    verdict = result.get("内容结论", "")
    explanation = result.get("内容说明") or ""
    if rule.get("字符类型") == "经纬度" and evidence_ok and evidence:
        # 坐标由代码精确换算比较，不依赖模型算术
        verdict_code = compare_coordinates(row["内容"], evidence)
        if verdict_code == "一致":
            verdict = "支持"
        elif verdict_code == "经纬度写反":
            verdict, explanation = "不支持", "经度和纬度写反了（按原文换算后比较）"
        elif verdict_code == "不一致":
            verdict = "不支持"
            explanation = explanation or "按原文换算后，相差超过 1′"
    if verdict == "不支持":
        issues.append(Issue("内容", "内容与资料不符", explanation + note,
                            建议值=result.get("建议填写内容"), 证据=evidence,
                            状态="错误" if sure and not is_descriptive(rule, row["内容"]) else "需复核"))
    elif verdict == "未提及":
        issues.append(Issue("来源", "资料不能证明该内容",
                            f"{doc.path.name} 中没有找到该数据。" + (result.get("内容说明") or ""),
                            证据=evidence, 状态="需复核"))
    elif verdict == "部分支持":
        issues.append(Issue("内容", "资料仅部分支持", (result.get("内容说明") or "") + note,
                            建议值=result.get("建议填写内容"), 证据=evidence, 状态="需复核"))

    rule_problem = result.get("规则问题")
    if rule_problem and not known + issues:
        issues.append(Issue("内容", "不符合填写要求", rule_problem, 状态="需复核"))
    return issues


def audit_row(row: dict, rules: Rules, kb: KnowledgeBase) -> list[Issue]:
    rule = effective_rule(row, rules)
    issues = hard_check(row, rule)
    needs_source = rule.get("需核查出处", "是") == "是"
    if needs_source and row["内容"] and row["来源"]:
        try:
            found = verify_with_source(row, rule, rules, kb, issues)
        except Exception as e:  # 单行失败不影响其他行
            found = [Issue("内容", "核查失败", f"核查出错：{e}", 状态="需复核")]
        # 模型按资料给出了修正值时，修正值已满足的格式问题就不再重复报
        fix = next((i for i in found if i.字段 == "内容" and i.状态 == "错误" and i.建议值), None)
        if fix:
            remaining = {i.问题类型 for i in hard_check({**row, "内容": fix.建议值}, rule) if i.字段 == "内容"}
            issues = [i for i in issues if i.字段 != "内容" or i.问题类型 in remaining]
        # 来源日期已按资料给出修正值时，日期的格式问题随之解决，不再重复报
        if any(i.字段 == "日期" and i.状态 == "错误" and i.建议值 for i in found):
            issues = [i for i in issues if not (i.字段 == "日期" and i.问题类型 == "格式错误")]
        issues += found
    return issues


def to_cell_value(s: str):
    """建议值是数字就按数字写回，保持和原表一致的类型。以 0 开头的编号（如邮编）保留为文本，以免丢掉前导 0。"""
    s = cell_str(s)
    num = to_number(s)
    if num is None or num != num or num in (float("inf"), float("-inf")):   # 认不出，或是 nan / inf
        return s
    if len(s) > 1 and s.startswith("0") and not s.startswith("0."):
        return s
    return int(num) if "." not in s else num


def write_outputs(path: Path, col_of: dict, header_row: int, rows: list, results: list) -> dict:
    config.OUTPUT_DIR.mkdir(exist_ok=True)
    wb = load_workbook(path)
    ws = wb.active
    result_col = ws.max_column + 1
    ws.cell(header_row, result_col, "审核结论").font = Font(bold=True)

    report = Workbook()
    rs = report.active
    rs.title = "错误报告"
    rs.append(["行号", cell_str(ws.cell(header_row, col_of["项目"]).value), "列", "原值", "问题类型", "处理", "说明", "建议值", "原文证据"])
    stats = {"文件": path.name, "总行数": len(rows), "通过": 0, "已自动修正": 0, "需人工复核": 0}

    for row, issues in zip(rows, results):
        r = row["_row"]
        fixed = review = 0
        for issue in issues:
            col = col_of.get(issue.字段)
            auto_fix = issue.状态 == "错误" and issue.建议值 not in (None, "", "null")
            if col:
                cell = ws.cell(r, col)
                old = cell_str(cell.value)
                if auto_fix:
                    cell.value = to_cell_value(issue.建议值)
                    cell.fill = FIXED_FILL
                    text = f"已修正\n原值：{old or '（空）'}\n原因：{issue.说明}"
                else:
                    cell.fill = REVIEW_FILL
                    text = f"需人工复核：{issue.问题类型}\n{issue.说明}"
                    if issue.建议值:
                        text += f"\n建议：{issue.建议值}"
                if issue.证据:
                    text += f"\n原文：{issue.证据}"
                prev = cell.comment.text + "\n---\n" if cell.comment else ""
                cell.comment = Comment(prev + text, "审核Agent", width=320, height=160)
            else:
                old = ""
            fixed += auto_fix
            review += not auto_fix
            rs.append([r, row["项目"], cell_str(ws.cell(header_row, col_of[issue.字段]).value) if col_of.get(issue.字段) else issue.字段, old or row.get(issue.字段, ""),
                       issue.问题类型, "已自动修正" if auto_fix else "需人工复核", issue.说明,
                       issue.建议值 or "", issue.证据])
        if not issues:
            conclusion = "通过"
            stats["通过"] += 1
        else:
            conclusion = "；".join(filter(None, [f"已修正{fixed}处" if fixed else "",
                                                 f"需复核{review}处" if review else ""]))
            stats["已自动修正"] += bool(fixed)
            stats["需人工复核"] += bool(review)
        ws.cell(r, result_col, conclusion)

    for i, width in enumerate([6, 18, 12, 22, 18, 12, 50, 14, 50], start=1):
        rs.column_dimensions[chr(64 + i)].width = width
    for cells in rs.iter_rows(min_row=2):
        for c in cells:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for c in rs[1]:
        c.font = Font(bold=True)

    wb.save(config.OUTPUT_DIR / f"{path.stem}_修正版.xlsx")
    report.save(config.OUTPUT_DIR / f"{path.stem}_错误报告.xlsx")
    return stats


def main():
    args = sys.argv[1:]
    if "--no-llm" in args:
        config.USE_LLM = "off"
        args.remove("--no-llm")
    files = [Path(a) for a in args] or sorted(
        p for p in config.INPUT_DIR.glob("*.xlsx") if not p.name.startswith("~$"))
    if not files:
        print(f"没有找到要审核的 Excel，请放到 {config.INPUT_DIR}")
        return
    rules = load_rules()
    kb = KnowledgeBase()
    print(f"规则：{len(rules.items)} 个项目规则，{len(rules.general)} 条通用规则；资料库：{len(kb.docs)} 份资料")
    print("模式：" + ("程序优先，程序拿不准的交给大模型" if llm_available()
                     else "不用大模型，程序拿不准的交给人工复核"))

    all_stats = []
    for path in files:
        try:
            col_of, rows, header_row = read_rows(path, rules)
        except (ValueError, OSError) as e:      # 一个文件读不了不影响其他文件
            print(f"\n跳过 {path.name}：{e}")
            continue
        print(f"\n审核 {path.name}（{len(rows)} 行）...")
        with ThreadPoolExecutor(config.MAX_WORKERS) as pool:
            results = list(pool.map(lambda row: audit_row(row, rules, kb), rows))
        for row, issues in zip(rows, results):
            mark = "✓" if not issues else "✗"
            print(f"  {mark} 第{row['_row']}行 {row['主体']}·{row['项目']}" +
                  "".join(f"\n      - [{i.状态}] {i.问题类型}：{i.说明}" for i in issues))
        try:
            all_stats.append(write_outputs(path, col_of, header_row, rows, results))
        except PermissionError:
            print(f"  无法保存 {path.stem} 的结果：请先在 Excel 里关闭 输出/ 中的同名文件，再重新运行")

    if not all_stats:
        print("\n没有成功审核的文件。")
        return
    summary = Workbook()
    ss = summary.active
    ss.title = "汇总"
    ss.append(list(all_stats[0].keys()))
    for s in all_stats:
        ss.append(list(s.values()))
    summary.save(config.OUTPUT_DIR / "汇总.xlsx")
    print(f"\n完成。结果在 {config.OUTPUT_DIR}")


if __name__ == "__main__":
    main()
