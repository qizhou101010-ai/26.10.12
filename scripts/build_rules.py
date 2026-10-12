"""根据 空白模板 + 填写正确的样例，让大模型自动总结审核规则，生成 审核规则.xlsx 初稿。

用法：python build_rules.py
生成后请打开 审核规则.xlsx 检查、修改，确认无误后再运行 audit.py。
"""
import json
import shutil
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation

import config
from audit import read_rows
from llm import ask_json
from rules import CHAR_TYPES, RULE_FIELDS, Rules, load_rules

SYSTEM_PROMPT = "你是数据填报审核专家，擅长从填报模板和样例中总结出明确、可执行的审核规则。以 JSON 格式输出。"

USER_PROMPT = """下面是一份填报表的【空白模板】（每行写明了填报项目、数据要求和字符要求），以及若干份【填写正确的样例】。
请为每个填报项目总结审核规则，并总结适用于所有行的通用规则。

【空白模板】
{template}

【填写正确的样例】
{samples}

要求：
- 数据类型只能是：数值 / 文本 / 日期。
- 字符类型（文本类项目用）：纯汉字 / 纯数字 / 纯字母 / 字母和数字 / 手机号 / 邮箱 / 身份证号 / 统一社会信用代码 / 邮政编码 / 不限；模板没要求就留空。
- 字符数：正好6个写“6”，范围写“2~10”，最多写“≤50”，至少写“≥2”；模板没要求就留空。
- 正则表达式：只有上面表达不了的固定格式（如编号规则）才写，否则留空。
- 取值范围写成 “最小值~最大值”，给一个宽松但能发现数量级或单位错误的范围；无法判断就留空字符串。
- 小数位：要求保留几位小数就写数字，没有要求留空字符串。
- 需核查出处：该项需要用资料来源证明就写“是”，否则写“否”。
- 补充要求：模板里有、但上面几项表达不了的要求，用一句话写清；没有就留空字符串。
- 通用规则：每条一句话，比如资料来源和日期的填写格式。

输出 JSON：
{{"项目规则": [{{"填报项目": "", "必填": "是/否", "数据类型": "", "字符类型": "", "字符数": "",
              "单位": "", "取值范围": "", "小数位": "", "需核查出处": "是/否", "补充要求": "",
              "正则表达式": ""}}],
  "通用规则": ["..."]}}"""


def rows_as_text(rows) -> str:
    keep = ["项目", "要求", "字符", "内容", "来源", "日期"]
    return json.dumps([{k: r[k] for k in keep if r.get(k)} for r in rows], ensure_ascii=False, indent=1)


def save_rules(path, columns: dict, items: list, general: list):
    wb = Workbook()
    ws = wb.active
    ws.title = "项目规则"
    ws.append(RULE_FIELDS)
    for item in items:
        ws.append([str(item.get(f, "") or "") for f in RULE_FIELDS])
    for i, w in enumerate([20, 6, 10, 16, 8, 10, 14, 8, 12, 40, 20], start=1):
        ws.column_dimensions[chr(64 + i)].width = w
    # 下拉菜单，减少手写出错
    dropdowns = {"必填": "是,否", "数据类型": "数值,文本,日期", "需核查出处": "是,否",
                 "字符类型": ",".join(["不限", *CHAR_TYPES])}
    for name, options in dropdowns.items():
        col = chr(65 + RULE_FIELDS.index(name))
        dv = DataValidation(type="list", formula1=f'"{options}"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}500")
    ws.freeze_panes = "B2"

    ws2 = wb.create_sheet("通用规则")
    ws2.append(["规则（每行一条，用文字描述）"])
    for g in general:
        ws2.append([g])
    ws2.column_dimensions["A"].width = 70

    ws3 = wb.create_sheet("列设置")
    ws3.append(["角色（不要改）", "填报表中的列名（可写多个，用顿号分隔；列名不同时在这里添加）"])
    for role, names in columns.items():
        ws3.append([role, "、".join(names)])
    ws3.column_dimensions["A"].width = 16
    ws3.column_dimensions["B"].width = 50

    for sheet in wb.worksheets:
        for c in sheet[1]:
            c.font = Font(bold=True)
    wb.save(path)


def main():
    rules = load_rules() if config.RULES_FILE.exists() else Rules()
    templates = sorted(config.TEMPLATE_DIR.glob("*.xlsx"))
    samples = sorted(config.SAMPLE_DIR.glob("*.xlsx"))
    if not templates:
        raise SystemExit(f"请先把空白模板放到 {config.TEMPLATE_DIR}")

    _, template_rows, _ = read_rows(templates[0], rules)
    sample_text = "\n\n".join(f"样例 {p.name}：\n{rows_as_text(read_rows(p, rules)[1])}" for p in samples)
    print(f"读取模板 {templates[0].name}（{len(template_rows)} 项），样例 {len(samples)} 份，正在总结规则...")

    result = ask_json(SYSTEM_PROMPT, USER_PROMPT.format(
        template=rows_as_text(template_rows), samples=sample_text or "（无样例）"))

    if config.RULES_FILE.exists():
        backup = config.RULES_FILE.with_name(f"审核规则_备份_{datetime.now():%Y%m%d_%H%M%S}.xlsx")
        shutil.copy2(config.RULES_FILE, backup)
        print(f"旧规则已备份为 {backup.name}")
    save_rules(config.RULES_FILE, rules.columns, result.get("项目规则", []), result.get("通用规则", []))
    print(f"已生成 {config.RULES_FILE.name}：{len(result.get('项目规则', []))} 个项目规则，"
          f"{len(result.get('通用规则', []))} 条通用规则。请打开检查、修改后再运行 audit.py。")


if __name__ == "__main__":
    main()
