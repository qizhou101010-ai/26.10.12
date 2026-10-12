"""自建资料库：读取资料文件，按《书名》找到资料，再检索出与填报内容最相关的段落。

命令行用法：
    python kb.py 导入 <文件夹>   # 把文件夹里的资料（含子文件夹）复制进资料库
    python kb.py 列表            # 查看资料库里有哪些资料
"""
import difflib
import math
import pickle
import re
import shutil
import sys
from dataclasses import dataclass, field
from pathlib import Path

import jieba
from rank_bm25 import BM25Okapi

import config

jieba.setLogLevel(60)
SUPPORTED = {".txt", ".md", ".docx", ".pdf"}
INDEX_FILE = config.KB_DIR / ".index.pkl"


class BM25(BM25Okapi):
    """默认的 IDF 在词出现于一半以上段落时会变成 0 或负数，资料少或对象名出现频繁时会把关键名称忽略掉。
    改用始终为正的 IDF：log((N+1)/(n+0.5))。"""

    def _calc_idf(self, nd):
        for word, freq in nd.items():
            self.idf[word] = math.log((self.corpus_size + 1) / (freq + 0.5))


def read_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md"}:
        raw = path.read_bytes()
        for enc in ("utf-8-sig", "gb18030"):
            try:
                return raw.decode(enc)
            except UnicodeDecodeError:
                continue
        return raw.decode("utf-8", errors="ignore")
    if suffix == ".docx":
        import docx
        d = docx.Document(str(path))
        lines = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                lines.append(" | ".join(c.text.strip() for c in row.cells))
        return "\n".join(lines)
    if suffix == ".pdf":
        from pypdf import PdfReader
        return "\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)
    return ""


def norm_title(s: str) -> str:
    """统一书名写法，便于匹配：去掉书名号、空白，全角括号转半角。"""
    s = re.sub(r"[《》<>〈〉\s]", "", str(s))
    return s.replace("（", "(").replace("）", ")").lower()


def tokenize(s: str) -> list[str]:
    """jieba 分词 + 中文双字切分。分词偶尔切错（如“青石山海拔”→“青石/山海拔”），双字能兜底。"""
    words = [w for w in jieba.lcut(s) if w.strip() and not re.fullmatch(r"[\W_]+", w)]
    for run in re.findall(r"[一-鿿]{2,}", s):
        words += [run[i:i + 2] for i in range(len(run) - 1)]
    return words


def split_chunks(text: str, size: int = 300) -> list[str]:
    """按段落切块，每块约 size 字，避免把一句话切断。"""
    chunks, buf = [], ""
    for para in (p.strip() for p in text.splitlines()):
        if not para:
            continue
        if buf and len(buf) + len(para) > size:
            chunks.append(buf)
            buf = ""
        buf = f"{buf}\n{para}" if buf else para
    if buf:
        chunks.append(buf)
    return chunks or [""]


@dataclass
class Doc:
    path: Path
    text: str
    titles: set = field(default_factory=set)
    chunks: list = field(default_factory=list)
    bm25: BM25 = None


class KnowledgeBase:
    def __init__(self, folder: Path = config.KB_DIR):
        self.folder = folder
        self.docs: list[Doc] = []
        self._load()

    def _load(self):
        cache = {}
        if INDEX_FILE.exists():
            try:
                cache = pickle.loads(INDEX_FILE.read_bytes())
            except Exception:
                cache = {}
        new_cache = {}
        for path in sorted(self.folder.rglob("*")):
            if path.suffix.lower() not in SUPPORTED or path.name.startswith(("~$", ".")):
                continue
            key, mtime = str(path.relative_to(self.folder)), path.stat().st_mtime
            if key in cache and cache[key][0] == mtime:
                text = cache[key][1]
            else:
                text = read_text(path)
            new_cache[key] = (mtime, text)
            first_line = next((l.strip() for l in text.splitlines() if l.strip()), "")
            doc = Doc(path=path, text=text,
                      titles={norm_title(path.stem), norm_title(first_line)},
                      chunks=split_chunks(text))
            doc.bm25 = BM25([tokenize(c) or ["_"] for c in doc.chunks])
            self.docs.append(doc)
        self.folder.mkdir(parents=True, exist_ok=True)
        INDEX_FILE.write_bytes(pickle.dumps(new_cache))
        # 跨资料的总索引：填写时还不知道该看哪份资料
        self._all = [(doc, c) for doc in self.docs for c in doc.chunks]
        self._bm25_all = BM25([tokenize(c) or ["_"] for _, c in self._all]) if self._all else None

    def find(self, source_name: str) -> Doc | None:
        """按资料名称找文件：先精确匹配文件名或首行标题，再模糊匹配。"""
        target = norm_title(source_name)
        if not target:
            return None
        for doc in self.docs:
            if target in doc.titles:
                return doc
        best, best_score = None, 0.0
        for doc in self.docs:
            for t in doc.titles:
                score = difflib.SequenceMatcher(None, target, t).ratio()
                if score > best_score:
                    best, best_score = doc, score
        return best if best_score >= 0.8 else None

    @staticmethod
    def search(doc: Doc, query: str, k: int = 4) -> list[str]:
        """资料不长时直接返回全文；否则返回开头一段（通常含标题和日期）+ 与 query 最相关的 k 段。"""
        if len(doc.text) <= config.FULL_TEXT_CHARS:
            return list(doc.chunks)
        scores = doc.bm25.get_scores(tokenize(query))
        ranked = sorted(range(len(doc.chunks)), key=lambda i: scores[i], reverse=True)
        picked = [0] + [i for i in ranked if i != 0][:k]
        return [doc.chunks[i] for i in sorted(set(picked))]


def search_all(kb: "KnowledgeBase", query: str, k: int = 6) -> list[tuple[Doc, str]]:
    """在整个资料库里检索与 query 最相关的 k 段，返回 (资料, 段落)。"""
    if kb._bm25_all is None:
        return []
    scores = kb._bm25_all.get_scores(tokenize(query))
    ranked = sorted(range(len(kb._all)), key=lambda i: scores[i], reverse=True)
    return [kb._all[i] for i in ranked[:k] if scores[i] > 0]


def evidence_in_doc(evidence: str, doc: Doc) -> bool:
    """检查模型给出的“原文证据”是否真的出现在资料里（忽略空白和标点差异）。"""
    strip = lambda s: re.sub(r"[\s\W_]+", "", s)
    ev = strip(evidence)
    return bool(ev) and ev in strip(doc.text)


def _cli():
    if len(sys.argv) >= 3 and sys.argv[1] == "导入":
        src = Path(sys.argv[2])
        count = 0
        for path in src.rglob("*"):
            if path.suffix.lower() in SUPPORTED and not path.name.startswith("~$"):
                dest = config.KB_DIR / path.relative_to(src)
                dest.parent.mkdir(parents=True, exist_ok=True)
                shutil.copy2(path, dest)
                count += 1
        print(f"已导入 {count} 个文件到 {config.KB_DIR}")
    elif len(sys.argv) < 2 or sys.argv[1] != "列表":
        print(__doc__)
        return
    kb = KnowledgeBase()
    print(f"资料库共有 {len(kb.docs)} 份资料：")
    for doc in kb.docs:
        print(f"  - {doc.path.relative_to(config.KB_DIR)}（{len(doc.text)} 字）")


if __name__ == "__main__":
    _cli()
