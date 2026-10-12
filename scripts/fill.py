"""第 2 步：按资料库里的资料填写模板，每项注明《来源》和来源日期。资料里找不到的项留空，不凭记忆填。

程序优先：数字、坐标、名称、来源日期由程序直接从资料中提取；程序提取不了的（如描述性内容）
才交给大模型；不用大模型时留空，并附上相关原文方便人工填写。

用法：
    python fill.py 模板/喜马拉雅山_模板.xlsx
    python fill.py 模板/喜马拉雅山_模板.xlsx --no-llm     # 完全不用大模型
生成 待审核/<模板名>_AI填写.xlsx，然后运行 audit.py 做独立审核。
"""
import sys
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

import config
import program_check
from audit import llm_available, read_rows, to_cell_value
from kb import KnowledgeBase, evidence_in_doc, search_all
from llm import ask_json
from rules import SUBJECT_ITEMS, cell_str, effective_rule, load_rules

AI_FILL = PatternFill("solid", fgColor="DDEBF7")    # 浅蓝：AI 填写的单元格

SYSTEM_PROMPT = """你是严谨的资料整理员，只根据提供的资料片段填写数据。
规则：
1. 只能使用资料片段中的信息，不得凭自己的记忆或常识填写。片段里没有，就返回 null。
2. 数据必须是针对【对象】本身的，不能拿其他对象的数据。
3. 严格按数据要求和字符要求的单位、格式、字数填写；描述性内容用简洁的话概括原文，不添加原文没有的说法。
4. 原文依据必须从片段中逐字摘录。
5. 以 JSON 格式输出。"""

USER_PROMPT = """【对象】{subject}
【数据类型】{项目}
【数据要求】{要求}
【字符要求】{字符}
{content_hint}
【资料片段】（每段标明了所属资料，资料开头一段通常含发布日期）
{passages}

输出 JSON：
{{"填写内容": "字符串，片段中找不到则为 null",
  "资料": "所用资料的名称，必须是片段标注的资料之一，加书名号，例如《XXX》；找不到则为 null",
  "资料发布日期": "该资料的发布日期 YYYY-MM-DD 或 YYYY-MM，找不到则为 null",
  "原文依据": "逐字摘录支撑填写内容的原文，不超过100字",
  "说明": "一句话"}}"""


def program_fill(row: dict, rule: dict, kb: KnowledgeBase) -> dict | None:
    """程序从各份资料里提取这一项；各资料提取结果唯一且一致时返回填写结果，否则返回 None 或说明。"""
    found = []
    for doc in kb.docs:
        if row["主体"] and row["主体"] not in doc.text:
            continue
        if got := program_check.extract_value(row, rule, doc.text):
            found.append((doc, *got))
    values = {v for _, v, _ in found}
    if len(values) > 1:
        detail = "；".join(f"《{d.path.stem}》为 {v}" for d, v, _ in found)
        return {"说明": f"【程序提取】不同资料的数值不一致（{detail}），留空待人工确认"}
    if not found:
        return None
    doc, value, sent = found[0]
    date, _ = program_check.find_publish_date(doc.text)
    return {"内容": value, "来源": f"《{doc.path.stem}》",
            "日期": program_check.format_date(date) if date else "",
            "说明": f"【程序提取】依据：{sent[:120]}" + ("" if date else "（资料里没找到发布日期，请人工补填）")}


def fill_row(row: dict, rule: dict, kb: KnowledgeBase) -> dict:
    """程序优先；程序提取不了时交给大模型；不用大模型时留空并附相关原文。"""
    result = program_fill(row, rule, kb)
    if result is not None:
        return result
    if llm_available():
        try:
            return llm_fill_row(row, kb)
        except Exception as e:   # 模型调用失败（如余额不足）时退回人工
            note = f"大模型调用失败（{str(e)[:60]}），"
    else:
        note = ""
    related = "；".join(f"《{d.path.stem}》：{program_check.related_text(row, d.text, 120)}"
                        for d in kb.docs if row["主体"] and row["主体"] in d.text)
    related = f"相关原文 {related}" if related else "资料库中没有找到提到该对象的资料"
    return {"说明": f"{note}程序无法自动提取，留空待人工填写。{related}"}


def llm_fill_row(row: dict, kb: KnowledgeBase) -> dict:
    """交给大模型填写，返回 {内容, 来源, 日期, 说明}；找不到可靠依据时内容为空。"""
    subject = row["主体"]
    hits = search_all(kb, f"{subject} {row['项目']} {row['要求']}", k=6)
    if not hits:
        return {"说明": "资料库中没有找到相关资料，留空"}
    blocks, seen = [], set()
    for doc, chunk in hits:
        if doc.chunks[0] not in seen:                 # 补上资料开头一段，便于找发布日期
            seen.add(doc.chunks[0])
            blocks.append(f"[资料《{doc.path.stem}》开头]\n{doc.chunks[0]}")
        if chunk not in seen:
            seen.add(chunk)
            blocks.append(f"[资料《{doc.path.stem}》]\n{chunk}")
    is_name = row["项目"] in SUBJECT_ITEMS
    hint = (f"【已填内容】{row['内容']}（只需找出提到该名称的资料，填写内容原样返回）\n" if is_name else "")
    result = ask_json(SYSTEM_PROMPT, USER_PROMPT.format(
        subject=subject, content_hint=hint, passages="\n\n".join(blocks), **row))

    value = result.get("填写内容")
    doc = kb.find(result.get("资料") or "")
    evidence = (result.get("原文依据") or "").strip()
    if not value or doc is None:
        return {"说明": "资料中未找到该数据，留空。" + (result.get("说明") or "")}
    if not any(d is doc for d, _ in hits):
        return {"说明": f"模型引用的资料《{doc.path.stem}》不在检索结果中，留空待人工填写"}
    if not evidence_in_doc(evidence, doc):
        return {"说明": "模型给出的原文依据在资料中找不到，可能是编造，留空待人工填写"}
    return {"内容": row["内容"] if is_name else str(value), "来源": f"《{doc.path.stem}》",
            "日期": result.get("资料发布日期") or "",
            "说明": f"【模型填写】依据：{evidence}"}


def fill_file(path: Path, kb: KnowledgeBase, rules) -> Path:
    col_of, rows, header_row = read_rows(path, rules)
    if not rows or not rows[0]["主体"]:
        raise SystemExit(f"{path.name}：请先在“名称”那一行的数据内容里填上对象名称")
    print(f"填写 {path.name}（对象：{rows[0]['主体']}，{len(rows)} 项）...")
    with ThreadPoolExecutor(config.MAX_WORKERS) as pool:
        results = list(pool.map(lambda r: fill_row(r, effective_rule(r, rules), kb), rows))

    wb = load_workbook(path)
    ws = wb.active
    note_col = ws.max_column + 1
    ws.cell(header_row, note_col, "填写说明").font = Font(bold=True)
    ws.column_dimensions[ws.cell(header_row, note_col).column_letter].width = 50
    for row, res in zip(rows, results):
        r = row["_row"]
        for role in ("内容", "来源", "日期"):
            if res.get(role) and col_of.get(role):
                cell = ws.cell(r, col_of[role])
                cell.value = to_cell_value(res[role]) if role == "内容" else res[role]
                cell.fill = AI_FILL
        ws.cell(r, note_col, res["说明"]).alignment = Alignment(wrap_text=True, vertical="top")
        status = "已填" if res.get("内容") else "留空"
        print(f"  [{status}] {row['项目']}：{cell_str(res.get('内容')) or '—'}  {res.get('来源', '')}")

    out = config.INPUT_DIR / f"{path.stem.removesuffix('_模板')}_AI填写.xlsx"
    out.parent.mkdir(exist_ok=True)
    wb.save(out)
    return out


def main():
    args = sys.argv[1:]
    if "--no-llm" in args:
        config.USE_LLM = "off"
        args.remove("--no-llm")
    files = [Path(a) for a in args]
    if not files:
        print(__doc__)
        return
    kb = KnowledgeBase()
    rules = load_rules()
    for path in files:
        out = fill_file(path, kb, rules)
        print(f"已生成 {out}\n下一步：python audit.py \"{out}\"")


if __name__ == "__main__":
    main()
